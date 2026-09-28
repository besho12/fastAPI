from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from app.comparison_engine.group import GroupComparisonEngine
from app.comparison_engine.gemini import GeminiGateway
from app.discrepancy_excel import generate_discrepancy_excel
from app.discrepancy_report import (
    CompanyFinding,
    CriterionFinding,
    DiscrepancyCriterion,
    DutyCompanyStatus,
    DutyGapFinding,
    JobDescriptionDiscrepancyReport,
    Priority,
    analyze_discrepancies,
    build_discrepancy_prompt,
)
from app.pipeline import JobComparisonPipeline
from app.schemas import CompanyInfo, JobDescription, JobInformation


COMPANIES = ["LSURV2602", "LSURV2604", "LSURV2608", "LSURV2609"]


def _jobs():
    return [
        JobDescription(
            job_id=f"job-{code}",
            company=CompanyInfo(company_code=code),
            job_information=JobInformation(
                job_code="TCGNBFI010104-A",
                job_title=f"Title for {code}",
            ),
            raw_text=f"Source job description for {code}",
        )
        for code in COMPANIES
    ]


def _report():
    criteria = []
    for criterion in DiscrepancyCriterion:
        criteria.append(
            CriterionFinding(
                criterion=criterion,
                company_findings=[
                    CompanyFinding(
                        company_code=code,
                        summary=f"{criterion.value} summary for {code}",
                        evidence=[f"Evidence for {code}"],
                    )
                    for code in COMPANIES
                ],
                benchmark_summary=f"Benchmark for {criterion.value}",
                discrepancy_observation=f"Difference for {criterion.value}",
                peer_evidence_summary=f"Peer evidence for {criterion.value}",
                suggested_review=f"Review {criterion.value}",
                priority=Priority.MEDIUM,
            )
        )

    return JobDescriptionDiscrepancyReport(
        job_code="TCGNBFI010104-A",
        target_company_code=COMPANIES[0],
        reference_company_codes=COMPANIES[1:],
        overall_summary="Overall grounded summary.",
        criteria=criteria,
        duty_gaps=[
            DutyGapFinding(
                duty_theme="Team leadership / supervision",
                company_statuses=[
                    DutyCompanyStatus(
                        company_code=code,
                        status="Explicit",
                        evidence=f"Leadership evidence for {code}",
                    )
                    for code in COMPANIES
                ],
                benchmark_observation="Aligned across the group.",
                review_status="Aligned",
            )
        ],
    )


class _Gateway:
    def __init__(self):
        self.calls = []

    def generate_model(self, **kwargs):
        self.calls.append(kwargs)
        return _report()


def test_gemini_gateway_supports_the_shared_structured_model_contract():
    gateway = object.__new__(GeminiGateway)
    calls = []

    def generate_json(**kwargs):
        calls.append(kwargs)
        return _report().model_dump(mode="json")

    gateway.generate_json = generate_json
    result = gateway.generate_model(
        system_instruction="System",
        user_content="User",
        response_schema=JobDescriptionDiscrepancyReport,
        label="gemini_discrepancy_test",
    )

    assert isinstance(result, JobDescriptionDiscrepancyReport)
    assert result.target_company_code == "LSURV2602"
    assert calls[0]["response_schema"] is JobDescriptionDiscrepancyReport


def test_prompt_labels_target_and_all_selected_peers():
    jobs = _jobs()
    prompt = build_discrepancy_prompt(jobs[0], jobs[1:])

    assert "TARGET COMPANY: LSURV2602" in prompt
    assert "REFERENCE COMPANIES: LSURV2604, LSURV2608, LSURV2609" in prompt
    assert prompt.count("BEGIN JOB DESCRIPTION") == 4


def test_structured_analysis_contract_and_workbook_layout():
    jobs = _jobs()
    gateway = _Gateway()
    report = analyze_discrepancies(
        target_job=jobs[0], reference_jobs=jobs[1:], gateway=gateway
    )

    assert [item.criterion for item in report.criteria] == list(
        DiscrepancyCriterion
    )
    assert len(gateway.calls) == 1

    payload = generate_discrepancy_excel(report)
    workbook = load_workbook(BytesIO(payload))
    assert workbook.sheetnames == [
        "Executive Summary",
        "Criteria Comparison",
        "Duty Gap Matrix",
    ]

    summary = workbook["Executive Summary"]
    assert summary["B4"].value == "LSURV2602"
    assert summary["A9"].value == "Duties & Responsibilities"
    assert summary.freeze_panes == "A5"

    criteria = workbook["Criteria Comparison"]
    assert [criteria.cell(3, column).value for column in range(2, 6)] == COMPANIES
    assert criteria["A8"].value == "Duties & Responsibilities"

    duties = workbook["Duty Gap Matrix"]
    assert duties["A4"].value == "Team leadership / supervision"
    workbook.close()


def test_group_zip_can_include_discrepancy_workbook():
    jobs = _jobs()
    group = GroupComparisonEngine().run(jobs)
    pipeline = object.__new__(JobComparisonPipeline)
    pipeline._model_gateway = _Gateway()

    payload, filename = pipeline._generate_group_comparison_package(
        "TCGNBFI010104-A",
        group,
        include_discrepancy_report=True,
    )

    assert filename == "Company_job_description_discrepancy_report.zip"
    with ZipFile(BytesIO(payload)) as archive:
        assert (
            "Company_job_description_discrepancy_report.xlsx"
            in archive.namelist()
        )
        assert len(archive.namelist()) == 1
        workbook = load_workbook(
            BytesIO(
                archive.read(
                    "Company_job_description_discrepancy_report.xlsx"
                )
            ),
            read_only=True,
        )
        assert workbook.sheetnames == [
            "Executive Summary",
            "Criteria Comparison",
            "Duty Gap Matrix",
        ]
        workbook.close()


def test_group_zip_reuses_precomputed_discrepancy_report():
    jobs = _jobs()
    group = GroupComparisonEngine().run(jobs)
    gateway = _Gateway()
    pipeline = object.__new__(JobComparisonPipeline)
    pipeline._model_gateway = gateway

    payload, _ = pipeline._generate_group_comparison_package(
        "TCGNBFI010104-A",
        group,
        include_discrepancy_report=True,
        discrepancy_report=_report(),
    )

    assert gateway.calls == []
    with ZipFile(BytesIO(payload)) as archive:
        assert (
            "Company_job_description_discrepancy_report.xlsx"
            in archive.namelist()
        )
        assert len(archive.namelist()) == 1
