"""
app/report_matrix_excel.py
"""

from __future__ import annotations

from io import BytesIO
from typing import List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas import JobDescription


# ==========================================================================
# Styling constants
# ==========================================================================

_HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
_HEADER_FONT = Font(
    bold=True,
    size=11,
    name="Aptos Narrow",
    color="FFFFFF",
)

_CRITERIA_FILL = PatternFill(fill_type="solid", fgColor="D9E1F2")
_CRITERIA_FONT = Font(
    bold=True,
    size=11,
    name="Aptos Narrow",
)

_VALUE_FONT = Font(
    bold=False,
    size=11,
    name="Aptos Narrow",
)

_LEFT_WRAP = Alignment(
    horizontal="left",
    vertical="top",
    wrap_text=True,
)

_CENTER_WRAP = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True,
)

_THIN_SIDE = Side(style="thin")

_THIN_BORDER = Border(
    left=_THIN_SIDE,
    right=_THIN_SIDE,
    top=_THIN_SIDE,
    bottom=_THIN_SIDE,
)

_NOT_SPECIFIED = "Not explicitly specified."


# ==========================================================================
# Criteria definitions
# ==========================================================================


def _job_code_text(job: JobDescription) -> str:
    job_code = job.job_information.job_code

    return (
        job_code.strip()
        if job_code and job_code.strip()
        else _NOT_SPECIFIED
    )


def _job_title_text(job: JobDescription) -> str:
    job_title = job.job_information.job_title

    return (
        job_title.strip()
        if job_title and job_title.strip()
        else _NOT_SPECIFIED
    )


def _bulleted(values: List[str]) -> str:
    cleaned = [
        str(v).strip()
        for v in values
        if v is not None and str(v).strip()
    ]

    if not cleaned:
        return _NOT_SPECIFIED

    if len(cleaned) == 1:
        return cleaned[0]

    return "\n".join(f"- {v}" for v in cleaned)


def _years_of_experience_text(job: JobDescription) -> str:
    return _bulleted(job.requirements.experience)

def _education_level_text(job: JobDescription) -> str:
    return _bulleted(job.requirements.education)

def _field_of_experience_text(job: JobDescription) -> str:
    return _bulleted(job.requirements.field_of_experience)


def _soft_skills_text(job: JobDescription) -> str:
    return _bulleted(job.requirements.soft_skills)


def _computer_skills_text(job: JobDescription) -> str:
    return _bulleted(job.requirements.computer)


def _language_requirements_text(job: JobDescription) -> str:
    return _bulleted(job.requirements.language)


def _duties_and_responsibilities_text(job: JobDescription) -> str:
    return _bulleted(job.responsibilities)


# ==========================================================================
# Criteria rows
# ==========================================================================

_CRITERIA_ROWS = [
    ("Job Code", _job_code_text),
    ("Job Title", _job_title_text),
    ("Years of Experience", _years_of_experience_text),
    ("Field of Experience", _field_of_experience_text),
    ("Education Level", _education_level_text),
    ("Soft Skills", _soft_skills_text),
    ("Computer Skills", _computer_skills_text),
    ("Language Requirements", _language_requirements_text),
    ("Duties and Responsibilities", _duties_and_responsibilities_text),
]


# ==========================================================================
# Column headers
# ==========================================================================


def _column_headers(
    new_job: JobDescription,
    reference_jobs: List[JobDescription],
) -> List[str]:
    all_jobs = [new_job] + reference_jobs

    headers = []

    for job in all_jobs:
        company_code = getattr(job.company, "company_code", None)

        if company_code and company_code.strip():
            headers.append(company_code.strip())
        else:
            headers.append(_NOT_SPECIFIED)

    return headers


# ==========================================================================
# Public API
# ==========================================================================


def generate_comparison_matrix_excel(
    new_job: JobDescription,
    reference_jobs: List[JobDescription],
) -> bytes:
    """
    Generate the "Job Comparison Matrix" Excel report entirely in memory.

    No .xlsx file is written to disk.

    Args:
        new_job:
            The job currently being evaluated.

        reference_jobs:
            All historical reference jobs retrieved for new_job's
            job code. Any count is supported, including zero.

    Returns:
        bytes:
            The generated .xlsx file as raw bytes.

    Raises:
        ValueError:
            On invalid None inputs.

        TypeError:
            On invalid input types.

        RuntimeError:
            If the generated workbook fails a basic sanity check.
    """

    # ==================================================================
    # Validation
    # ==================================================================

    if new_job is None:
        raise ValueError("new_job must not be None")

    if not isinstance(new_job, JobDescription):
        raise TypeError(
            "new_job must be a JobDescription instance"
        )

    if reference_jobs is None:
        raise ValueError(
            "reference_jobs must not be None"
        )

    if not isinstance(reference_jobs, list):
        raise TypeError(
            "reference_jobs must be a list"
        )

    for job in reference_jobs:
        if not isinstance(job, JobDescription):
            raise TypeError(
                "Every reference job must be a JobDescription instance"
            )

    # ==================================================================
    # Create workbook
    # ==================================================================

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Job Comparison Matrix"

    all_jobs = [new_job] + reference_jobs

    headers = _column_headers(
        new_job,
        reference_jobs,
    )

    # ==================================================================
    # Header row (row 1)
    # ==================================================================

    worksheet.cell(
        row=1,
        column=1,
        value="Criteria",
    )

    worksheet.cell(
        row=1,
        column=1,
    ).font = _HEADER_FONT

    worksheet.cell(
        row=1,
        column=1,
    ).fill = _HEADER_FILL

    worksheet.cell(
        row=1,
        column=1,
    ).alignment = _CENTER_WRAP

    for col_index, header_label in enumerate(
        headers,
        start=2,
    ):
        cell = worksheet.cell(
            row=1,
            column=col_index,
            value=header_label,
        )

        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER_WRAP

    # ==================================================================
    # Criteria rows
    # ==================================================================

    for row_offset, (
        label,
        extractor,
    ) in enumerate(
        _CRITERIA_ROWS,
        start=2,
    ):
        label_cell = worksheet.cell(
            row=row_offset,
            column=1,
            value=label,
        )

        label_cell.font = _CRITERIA_FONT
        label_cell.fill = _CRITERIA_FILL
        label_cell.alignment = _CENTER_WRAP

        for col_offset, job in enumerate(
            all_jobs,
            start=2,
        ):
            value = extractor(job)

            value_cell = worksheet.cell(
                row=row_offset,
                column=col_offset,
                value=value,
            )

            value_cell.font = _VALUE_FONT
            value_cell.alignment = _LEFT_WRAP

    # ==================================================================
    # Borders on every used cell
    # ==================================================================

    total_rows = 1 + len(_CRITERIA_ROWS)
    total_cols = 1 + len(all_jobs)

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=total_rows,
        min_col=1,
        max_col=total_cols,
    ):
        for cell in row:
            cell.border = _THIN_BORDER

    # ==================================================================
    # Column widths
    # ==================================================================

    worksheet.column_dimensions["A"].width = 26

    for col_index in range(
        2,
        total_cols + 1,
    ):
        column_letter = get_column_letter(
            col_index
        )

        worksheet.column_dimensions[
            column_letter
        ].width = 42

    # ==================================================================
    # Row heights
    # ==================================================================

    worksheet.row_dimensions[1].height = 30

    for row_offset in range(
        2,
        total_rows + 1,
    ):
        max_line_count = 1

        for col_offset in range(
            1,
            total_cols + 1,
        ):
            content = worksheet.cell(
                row=row_offset,
                column=col_offset,
            ).value

            if content:
                max_line_count = max(
                    max_line_count,
                    str(content).count("\n") + 1,
                )

        worksheet.row_dimensions[
            row_offset
        ].height = max(
            24,
            min(
                220,
                16 * max_line_count,
            ),
        )

    # ==================================================================
    # Freeze header row + criteria column
    # ==================================================================

    worksheet.freeze_panes = "B2"

    # ==================================================================
    # Save workbook IN MEMORY
    # ==================================================================

    excel_buffer = BytesIO()

    workbook.save(excel_buffer)

    excel_buffer.seek(0)

    excel_bytes = excel_buffer.getvalue()

    # Close the workbook after serialization.
    workbook.close()

    # ==================================================================
    # Basic sanity check
    #
    # IMPORTANT:
    # We no longer check Path.exists() / file size because the workbook
    # is intentionally never written to disk.
    # ==================================================================

    if not excel_bytes:
        raise RuntimeError(
            "Comparison matrix Excel generation failed: "
            "generated workbook is empty"
        )

    # ==================================================================
    # Verify generated workbook IN MEMORY
    # ==================================================================

    verification_buffer = BytesIO(excel_bytes)

    verification_workbook = load_workbook(
        verification_buffer,
        read_only=True,
        data_only=True,
    )

    try:
        # --------------------------------------------------------------
        # Worksheet check
        # --------------------------------------------------------------

        if (
            "Job Comparison Matrix"
            not in verification_workbook.sheetnames
        ):
            raise RuntimeError(
                "Generated comparison matrix Excel does not contain "
                "'Job Comparison Matrix' worksheet"
            )

        verification_sheet = (
            verification_workbook[
                "Job Comparison Matrix"
            ]
        )

        # --------------------------------------------------------------
        # A1 check
        # --------------------------------------------------------------

        if verification_sheet["A1"].value != "Criteria":
            raise RuntimeError(
                "Generated comparison matrix Excel contract violation: "
                "A1 must be 'Criteria'"
            )

        # --------------------------------------------------------------
        # B1 check
        # --------------------------------------------------------------

        expected_company_code = getattr(
            new_job.company,
            "company_code",
            None,
        )

        if expected_company_code:
            expected_company_code = (
                expected_company_code.strip()
            )
        else:
            expected_company_code = _NOT_SPECIFIED

        if (
            verification_sheet["B1"].value
            != expected_company_code
        ):
            raise RuntimeError(
                "Generated comparison matrix Excel contract violation: "
                f"B1 must be '{expected_company_code}'"
            )

        # --------------------------------------------------------------
        # A2 check
        # --------------------------------------------------------------

        if verification_sheet["A2"].value != "Job Code":
            raise RuntimeError(
                "Generated comparison matrix Excel contract violation: "
                "A2 must be 'Job Code'"
            )

        # --------------------------------------------------------------
        # A3 check
        # --------------------------------------------------------------

        if verification_sheet["A3"].value != "Job Title":
            raise RuntimeError(
                "Generated comparison matrix Excel contract violation: "
                "A3 must be 'Job Title'"
            )

    finally:
        verification_workbook.close()
        verification_buffer.close()

    # ==================================================================
    # Return raw Excel bytes
    # ==================================================================

    return excel_bytes


# ==========================================================================
# Public exports
# ==========================================================================

__all__ = [
    "generate_comparison_matrix_excel",
]