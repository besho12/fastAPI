"""Deterministic Excel renderer for the target-company discrepancy report."""

from __future__ import annotations

from io import BytesIO
from typing import Dict, Iterable, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.discrepancy_report import (
    CompanyFinding,
    JobDescriptionDiscrepancyReport,
)


NAVY = "17365D"
LIGHT_BLUE = "D9EAF7"
NOTE_FILL = "F3F6F9"
WHITE = "FFFFFF"
GREY = "666666"
MAX_CELL_TEXT = 32767


def _text(value: object) -> str:
    return str(value or "").strip()[:MAX_CELL_TEXT]


def _header(ws, row: int, values: Iterable[str]) -> None:
    for column, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Carlito", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )


def _body(cell, *, fill: str | None = None) -> None:
    cell.font = Font(name="Carlito", size=11)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def _title(ws, text: str, last_column: int, subtitle: bool = False) -> None:
    row = 2 if subtitle else 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(
        name="Carlito",
        size=11 if subtitle else 16,
        bold=not subtitle,
        color=GREY if subtitle else WHITE,
    )
    cell.fill = PatternFill("solid", fgColor="000000" if subtitle else NAVY)
    if subtitle:
        cell.fill = PatternFill(fill_type=None)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 24


def _company_map(findings: List[CompanyFinding]) -> Dict[str, CompanyFinding]:
    return {item.company_code.strip().upper(): item for item in findings}


def _finding_text(finding: CompanyFinding | None, include_evidence: bool = False) -> str:
    if finding is None:
        return "Not explicitly specified."
    value = _text(finding.summary) or "Not explicitly specified."
    if include_evidence and finding.evidence:
        evidence = "\n".join(f"- {_text(item)}" for item in finding.evidence)
        value = f"{value}\n\nSource evidence:\n{evidence}"
    return value


def _set_widths(ws, widths: Iterable[float]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _executive_summary(report: JobDescriptionDiscrepancyReport):
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    _title(ws, "Main Company Job Description Benchmark Report", 7)
    _title(
        ws,
        (
            f"Benchmark of {report.target_company_code} against "
            f"{', '.join(report.reference_company_codes)} | "
            f"Same Job Code: {report.job_code}"
        ),
        7,
        subtitle=True,
    )
    _header(
        ws,
        4,
        [
            "Criterion",
            report.target_company_code,
            "Peer Benchmark",
            "Discrepancy / Observation",
            "Peer Evidence",
            f"Suggested Review for {report.target_company_code}",
            "Priority",
        ],
    )

    for row, criterion in enumerate(report.criteria, start=5):
        findings = _company_map(criterion.company_findings)
        values = [
            criterion.criterion.value,
            _finding_text(findings.get(report.target_company_code), True),
            criterion.benchmark_summary,
            criterion.discrepancy_observation,
            criterion.peer_evidence_summary,
            criterion.suggested_review,
            criterion.priority.value,
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=column, value=_text(value))
            _body(cell, fill=LIGHT_BLUE if column == 2 else None)

    summary_row = 5 + len(report.criteria) + 1
    ws.merge_cells(
        start_row=summary_row, start_column=1, end_row=summary_row, end_column=7
    )
    overall = ws.cell(
        summary_row,
        1,
        f"Overall assessment: {_text(report.overall_summary)}",
    )
    _body(overall, fill=LIGHT_BLUE)
    overall.font = Font(name="Carlito", size=11, bold=True)

    note_row = summary_row + 1
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    note = ws.cell(
        note_row,
        1,
        (
            "Note: A gap means content is explicit in one or more peer job "
            f"descriptions but not explicit in {report.target_company_code}. "
            "It does not automatically mean the target company should adopt "
            "it; organisational mandate and authority must be validated."
        ),
    )
    _body(note, fill=NOTE_FILL)
    note.font = Font(name="Carlito", size=11, color=GREY)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:G{4 + len(report.criteria)}"
    _set_widths(ws, [22, 38, 48, 54, 44, 52, 13])
    return wb


def _criteria_sheet(wb: Workbook, report: JobDescriptionDiscrepancyReport) -> None:
    ws = wb.create_sheet("Criteria Comparison")
    companies = [report.target_company_code, *report.reference_company_codes]
    last_column = len(companies) + 2
    _title(ws, "Side-by-Side Source Comparison", last_column)
    headers = ["Criterion", *companies, "Key Target-Company Difference"]
    _header(ws, 3, headers)

    for row, criterion in enumerate(report.criteria, start=4):
        findings = _company_map(criterion.company_findings)
        values = [
            criterion.criterion.value,
            *[_finding_text(findings.get(code), True) for code in companies],
            criterion.discrepancy_observation,
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row, column, _text(value))
            _body(cell, fill=LIGHT_BLUE if column == 2 else None)

    note_row = 4 + len(report.criteria) + 1
    ws.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row,
        end_column=last_column,
    )
    note = ws.cell(
        note_row,
        1,
        "Source basis: the job descriptions stored for the selected companies. "
        "Not explicitly stated is not treated as confirmed absence.",
    )
    _body(note, fill=NOTE_FILL)
    note.font = Font(name="Carlito", size=11, color=GREY)
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:{get_column_letter(last_column)}{3 + len(report.criteria)}"
    _set_widths(ws, [24, *([44] * len(companies)), 56])


def _duty_sheet(wb: Workbook, report: JobDescriptionDiscrepancyReport) -> None:
    ws = wb.create_sheet("Duty Gap Matrix")
    companies = [report.target_company_code, *report.reference_company_codes]
    last_column = len(companies) + 3
    _title(ws, "Duties & Responsibilities - Benchmark Gap Matrix", last_column)
    _header(
        ws,
        3,
        [
            "Duty Theme",
            *companies,
            "Benchmark Observation",
            "Review Status",
        ],
    )

    for row, gap in enumerate(report.duty_gaps, start=4):
        statuses = {
            item.company_code.strip().upper(): item for item in gap.company_statuses
        }
        company_values = []
        for code in companies:
            item = statuses.get(code)
            if item is None:
                company_values.append("Not explicitly specified.")
            else:
                evidence = f"\nEvidence: {_text(item.evidence)}" if item.evidence else ""
                company_values.append(f"{_text(item.status)}{evidence}")
        values = [
            gap.duty_theme,
            *company_values,
            gap.benchmark_observation,
            gap.review_status,
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row, column, _text(value))
            _body(cell, fill=LIGHT_BLUE if column == 2 else None)

    note_row = 4 + len(report.duty_gaps) + 1
    ws.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row,
        end_column=last_column,
    )
    note = ws.cell(
        note_row,
        1,
        (
            "Interpretation: Gap to review identifies peer duties not explicit "
            "in the target description; Scope decision identifies duties that "
            "may intentionally sit elsewhere in the organisation."
        ),
    )
    _body(note, fill=NOTE_FILL)
    note.font = Font(name="Carlito", size=11, color=GREY)
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:{get_column_letter(last_column)}{3 + len(report.duty_gaps)}"
    _set_widths(ws, [30, *([35] * len(companies)), 52, 18])


def generate_discrepancy_excel(
    report: JobDescriptionDiscrepancyReport,
) -> bytes:
    wb = _executive_summary(report)
    _criteria_sheet(wb, report)
    _duty_sheet(wb, report)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = BytesIO()
    wb.save(buffer)
    payload = buffer.getvalue()
    if not payload:
        raise RuntimeError("Discrepancy workbook generation returned no data.")

    check = load_workbook(BytesIO(payload), read_only=True)
    if check.sheetnames != [
        "Executive Summary",
        "Criteria Comparison",
        "Duty Gap Matrix",
    ]:
        raise RuntimeError("Discrepancy workbook sheet contract failed.")
    check.close()
    return payload


__all__ = ["generate_discrepancy_excel"]
