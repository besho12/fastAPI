"""Excel overview for an all-jobs job-code comparison."""

from __future__ import annotations

from io import BytesIO
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.comparison_engine.group import JobGroupComparison
from app.comparison_engine.models import (
    BUCKET_LABEL,
    CATEGORY_ORDER,
    Bucket,
    Direction,
)


NAVY = "1F3864"
LIGHT_BLUE = "D9EAF7"
LIGHT_RED = "FCE4E4"
LIGHT_PURPLE = "F0E6F6"
LIGHT_GREEN = "E2F0D9"


def _header(sheet, row: int, labels: Iterable[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )


def _percent(value: Optional[float]) -> str:
    return "Not measurable" if value is None else f"{value * 100:.0f}%"


def _fit(sheet, widths) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _direction_summary(report, direction: Direction) -> tuple[str, int]:
    """Render one readable, category-grouped direction into a single cell."""

    findings = [
        finding
        for finding in report.findings
        if finding.direction == direction
    ]

    if not findings:
        message = (
            "No missing requirements compared with the other jobs."
            if direction == Direction.M_MINUS
            else "No additional requirements compared with the other jobs."
        )
        return message, 1

    lines = []
    for category in CATEGORY_ORDER:
        category_findings = [
            finding
            for finding in findings
            if finding.category == category
        ]
        if not category_findings:
            continue

        lines.append(category.value.upper())
        for finding in category_findings:
            benchmark_prefix = (
                "seen in"
                if direction == Direction.M_MINUS
                else "also seen in"
            )
            benchmark = (
                f"{benchmark_prefix} {finding.support_count} of "
                f"{finding.applicable_reference_count} other jobs"
                if finding.applicable_reference_count
                else "benchmark availability not measurable"
            )
            priority = BUCKET_LABEL[finding.bucket]
            lines.append(f"   • {finding.label}")
            lines.append(f"      {priority} · {benchmark}")

    # Excel cells have a 32,767-character limit.  Extremely large groups
    # retain the complete row-level detail in "Per-job Differences".
    text = "\n".join(lines)
    if len(text) > 32_000:
        text = (
            text[:31_900].rsplit("\n", 1)[0]
            + "\n… More items are available in Per-job Differences."
        )

    return text, text.count("\n") + 1


def generate_group_summary_excel(group: JobGroupComparison) -> bytes:
    """Create a compact decision-first workbook for the whole group."""

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Group Overview"
    overview.sheet_view.showGridLines = False

    _header(
        overview,
        1,
        [
            "Company",
            "Job title",
            "Job code",
            "Compared with",
            "Alignment",
            "Missing requirements",
            "Additional requirements",
            "Critical gaps",
            "Notable gaps",
            "Confidence",
        ],
    )

    for row, report in enumerate(group.reports, start=2):
        missing = sum(
            1 for finding in report.findings
            if finding.direction == Direction.M_MINUS
        )
        additional = sum(
            1 for finding in report.findings
            if finding.direction == Direction.M_PLUS
        )
        values = [
            report.new_company_code or "Not specified",
            report.new_job_title or "Not specified",
            report.new_job_code or "Not specified",
            report.reference_jobs_count,
            _percent(report.alignment_score),
            missing,
            additional,
            report.count_of(Bucket.CRITICAL_GAP),
            report.count_of(Bucket.NOTABLE_GAP),
            _percent(report.confidence),
        ]
        for column, value in enumerate(values, start=1):
            cell = overview.cell(row=row, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if column == 6 and missing:
                cell.fill = PatternFill("solid", fgColor=LIGHT_RED)
            elif column == 7 and additional:
                cell.fill = PatternFill("solid", fgColor=LIGHT_PURPLE)

    _fit(overview, [18, 28, 18, 14, 14, 20, 21, 14, 14, 14])

    company_summary = workbook.create_sheet("Company Comparison")
    company_summary.sheet_view.showGridLines = False
    _header(
        company_summary,
        1,
        [
            "Company",
            "Job title",
            "Job code",
            "Compared with",
            "Missing count",
            "Missing compared with other companies",
            "Additional count",
            "Additional compared with other companies",
        ],
    )
    company_summary.row_dimensions[1].height = 32

    for row, report in enumerate(group.reports, start=2):
        missing_summary, missing_lines = _direction_summary(
            report,
            Direction.M_MINUS,
        )
        additional_summary, additional_lines = _direction_summary(
            report,
            Direction.M_PLUS,
        )
        missing_count = sum(
            1
            for finding in report.findings
            if finding.direction == Direction.M_MINUS
        )
        additional_count = sum(
            1
            for finding in report.findings
            if finding.direction == Direction.M_PLUS
        )
        values = [
            report.new_company_code or "Not specified",
            report.new_job_title or "Not specified",
            report.new_job_code or "Not specified",
            report.reference_jobs_count,
            missing_count,
            missing_summary,
            additional_count,
            additional_summary,
        ]

        for column, value in enumerate(values, start=1):
            cell = company_summary.cell(row=row, column=column, value=value)
            cell.alignment = Alignment(
                vertical="top",
                horizontal="left",
                wrap_text=True,
            )

        company_summary.cell(row=row, column=1).font = Font(
            name="Calibri",
            bold=True,
            color=NAVY,
        )
        company_summary.cell(row=row, column=5).alignment = Alignment(
            horizontal="center",
            vertical="top",
        )
        company_summary.cell(row=row, column=6).fill = PatternFill(
            "solid",
            fgColor=LIGHT_RED if missing_count else LIGHT_GREEN,
        )
        company_summary.cell(row=row, column=7).alignment = Alignment(
            horizontal="center",
            vertical="top",
        )
        company_summary.cell(row=row, column=8).fill = PatternFill(
            "solid",
            fgColor=LIGHT_PURPLE if additional_count else LIGHT_BLUE,
        )
        company_summary.row_dimensions[row].height = min(
            360,
            max(36, max(missing_lines, additional_lines) * 15),
        )

    _fit(company_summary, [18, 30, 18, 16, 15, 78, 16, 78])

    differences = workbook.create_sheet("Per-job Differences")
    differences.sheet_view.showGridLines = False
    _header(
        differences,
        1,
        [
            "Company",
            "Category",
            "Difference",
            "Requirement",
            "Other jobs containing it",
            "Other applicable jobs",
            "Prevalence",
            "Priority",
            "Recommended action",
        ],
    )
    row = 2
    for report in group.reports:
        # Keep all missing rows together, followed by all additional rows.
        # The engine's importance order remains intact within each block.
        ordered_findings = [
            finding
            for direction in (Direction.M_MINUS, Direction.M_PLUS)
            for finding in report.findings
            if finding.direction == direction
        ]
        for finding in ordered_findings:
            difference = (
                "Missing from this job"
                if finding.direction == Direction.M_MINUS
                else "Additional in this job"
            )
            values = [
                report.new_company_code or "Not specified",
                finding.category.value,
                difference,
                finding.label,
                finding.support_count,
                finding.applicable_reference_count,
                finding.prevalence_text,
                finding.bucket.value.replace("_", " ").title(),
                finding.action_text,
            ]
            for column, value in enumerate(values, start=1):
                cell = differences.cell(row=row, column=column, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.fill = PatternFill(
                    "solid",
                    fgColor=(
                        LIGHT_RED
                        if finding.direction == Direction.M_MINUS
                        else LIGHT_PURPLE
                    ),
                )
                if column == 3:
                    cell.font = Font(
                        name="Calibri",
                        bold=True,
                        color=(
                            "9C0006"
                            if finding.direction == Direction.M_MINUS
                            else "7030A0"
                        ),
                    )
            row += 1
    if row == 2:
        differences.cell(row=2, column=1, value="No differences found.")
    _fit(differences, [18, 24, 23, 42, 20, 20, 17, 20, 28])

    matrix = workbook.create_sheet("Concept Coverage")
    matrix.sheet_view.showGridLines = False
    companies = [
        report.new_company_code or f"Job {index}"
        for index, report in enumerate(group.reports, start=1)
    ]
    _header(matrix, 1, ["Category", "Concept", *companies])

    evidence_by_report = [
        {evidence.concept_id: evidence for evidence in report.evidence}
        for report in group.reports
    ]
    concept_order = []
    seen = set()
    for report in group.reports:
        for evidence in report.evidence:
            if evidence.concept_id not in seen:
                seen.add(evidence.concept_id)
                concept_order.append(evidence)

    for row, concept in enumerate(concept_order, start=2):
        matrix.cell(row=row, column=1, value=concept.category.value)
        matrix.cell(row=row, column=2, value=concept.label)
        for report_index, lookup in enumerate(evidence_by_report, start=3):
            evidence = lookup.get(concept.concept_id)
            present = bool(evidence and evidence.new_job_has)
            cell = matrix.cell(
                row=row,
                column=report_index,
                value="Present" if present else "Missing",
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = PatternFill(
                "solid", fgColor=LIGHT_BLUE if present else LIGHT_RED
            )

    _fit(matrix, [25, 45, *([18] * len(companies))])

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    payload = buffer.getvalue()
    buffer.close()
    if not payload:
        raise RuntimeError("Group Excel generation produced no data.")
    return payload


__all__ = ["generate_group_summary_excel"]
