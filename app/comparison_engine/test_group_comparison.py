"""Tests for the all-jobs-against-all-others workflow."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO
from types import SimpleNamespace
from typing import List
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.comparison_engine.group import GroupComparisonEngine
from app.comparison_engine.group_excel import generate_group_summary_excel
from app.comparison_engine.models import Direction
from app.pipeline import JobComparisonPipeline
from app.schemas import CompanyInfo, JobDescription, JobInformation, Requirements


@dataclass
class _Company:
    company_code: str


@dataclass
class _JobInformation:
    job_code: str = "CLERK-01"
    job_title: str = "Office Clerk"


@dataclass
class _Requirements:
    experience: List[str] = field(default_factory=list)
    field_of_experience: List[str] = field(default_factory=list)
    education: List[str] = field(default_factory=list)
    language: List[str] = field(default_factory=list)
    computer: List[str] = field(default_factory=list)
    soft_skills: List[str] = field(default_factory=list)


@dataclass
class _Job:
    job_id: str
    company: _Company
    job_information: _JobInformation = field(default_factory=_JobInformation)
    requirements: _Requirements = field(default_factory=_Requirements)
    responsibilities: List[str] = field(default_factory=list)


def _job(company: str, *, language=(), computer=()) -> _Job:
    return _Job(
        job_id=f"job-{company}",
        company=_Company(company),
        requirements=_Requirements(
            language=list(language),
            computer=list(computer),
        ),
    )


def _three_jobs():
    return [
        _job("A", language=["Arabic"], computer=["Typing Speed"]),
        _job("B", language=["Native English"], computer=["Excel"]),
        _job("C", language=["Native English"], computer=["Excel"]),
    ]


def _finding(report, label: str):
    return next(
        finding
        for finding in report.findings
        if finding.label.casefold() == label.casefold()
    )


def test_each_job_gets_its_own_all_others_report():
    group = GroupComparisonEngine().run(_three_jobs())

    assert len(group.reports) == 3
    assert all(report.reference_jobs_count == 2 for report in group.reports)
    assert [report.new_company_code for report in group.reports] == [
        "A",
        "B",
        "C",
    ]

    report_a = group.reports[0]
    english = _finding(report_a, "Native English")
    typing = _finding(report_a, "Typing Speed")

    assert english.direction == Direction.M_MINUS
    assert english.support_count == 2
    assert english.applicable_reference_count == 2
    assert typing.direction == Direction.M_PLUS
    assert typing.support_count == 0

    # The same shared concept is recalculated from B's perspective.
    report_b = group.reports[1]
    typing_for_b = _finding(report_b, "Typing Speed")
    assert typing_for_b.direction == Direction.M_MINUS
    assert typing_for_b.support_count == 1
    assert typing_for_b.applicable_reference_count == 2


class _CountingGateway:
    def __init__(self):
        self.labels = []
        self.call_count = 0
        self.failure_count = 0

    def generate_json(
        self, system_instruction, user_content, response_schema, label=""
    ):
        self.labels.append(label)
        self.call_count += 1
        if label.startswith("clustering"):
            ids = list(
                dict.fromkeys(
                    re.findall(r"[A-Z]{3}-D\d{2}-\d{2}", user_content)
                )
            )
            return {
                "concepts": [
                    {"label": item_id, "member_item_ids": [item_id]}
                    for item_id in ids
                ]
            }
        return {"verdicts": []}


def test_semantic_clustering_runs_once_per_populated_category_for_group():
    gateway = _CountingGateway()
    GroupComparisonEngine(gateway=gateway).run(_three_jobs())

    clustering_calls = [
        label for label in gateway.labels if label.startswith("clustering")
    ]
    assert len(clustering_calls) == 2  # Language + Computer Skills
    assert len(set(clustering_calls)) == 2
    assert gateway.labels.count("materiality") == 1


def test_group_summary_workbook_is_decision_first():
    group = GroupComparisonEngine().run(_three_jobs())
    payload = generate_group_summary_excel(group)
    workbook = load_workbook(BytesIO(payload), read_only=True)

    assert workbook.sheetnames == [
        "Group Overview",
        "Per-job Differences",
        "Concept Coverage",
    ]
    overview = workbook["Group Overview"]
    assert overview.max_row == 4
    assert overview.cell(2, 1).value == "A"
    assert overview.cell(2, 6).value >= 1

    differences = workbook["Per-job Differences"]
    rows = list(differences.iter_rows(min_row=2, values_only=True))
    assert any(
        row[0] == "A"
        and row[2] == "Missing from this job"
        and row[3] == "Native English"
        for row in rows
    )
    workbook.close()


def test_pipeline_zip_contains_summary_and_one_workbook_per_job():
    group = GroupComparisonEngine().run(_three_jobs())
    pipeline = object.__new__(JobComparisonPipeline)

    payload, filename = pipeline._generate_group_comparison_package(
        "CLERK-01", group
    )

    assert filename.endswith("_all_jobs_comparison.zip")
    with ZipFile(BytesIO(payload)) as archive:
        names = archive.namelist()
        assert "CLERK-01_group_summary.xlsx" in names
        assert len([name for name in names if name.startswith("jobs/")]) == 3
        assert archive.testzip() is None


def test_existing_selected_job_method_now_returns_whole_group(
    monkeypatch,
):
    jobs = [
        JobDescription(
            job_id=job.job_id,
            company=CompanyInfo(company_code=job.company.company_code),
            job_information=JobInformation(
                job_code=job.job_information.job_code,
                job_title=job.job_information.job_title,
            ),
            requirements=Requirements(
                language=job.requirements.language,
                computer=job.requirements.computer,
            ),
            raw_text=f"Description for {job.company.company_code}",
        )
        for job in _three_jobs()
    ]
    rows = [
        SimpleNamespace(
            job=job,
            job_id=job.job_id,
            content_hash=f"content-{job.job_id}",
            raw_content_hash=f"raw-{job.job_id}",
        )
        for job in jobs
    ]

    class _Repository:
        def get_job_by_job_code_and_company_code(self, job_code, company_code):
            return next(
                row
                for row in rows
                if row.job.company.company_code == company_code
            )

        def get_jobs_by_job_code(self, job_code):
            return rows

    monkeypatch.setattr(
        "app.pipeline.job_row_to_job_description", lambda row: row.job
    )

    pipeline = object.__new__(JobComparisonPipeline)
    pipeline._job_repository = _Repository()
    pipeline._comparison_repository = None
    pipeline._report_repository = None
    pipeline._db_session = None
    pipeline._group_comparison_engine = GroupComparisonEngine()

    result = pipeline.compare_selected_job("clerk-01", "a")

    assert result.job_code == "CLERK-01"
    assert result.new_job.company.company_code == "A"
    assert len(result.group_reports) == 3
    assert result.report.new_company_code == "A"
    with ZipFile(BytesIO(result.zip_bytes)) as archive:
        assert len(archive.namelist()) == 4


def test_group_history_rows_are_persisted_for_every_job_atomically(
    monkeypatch,
):
    group = GroupComparisonEngine().run(_three_jobs())
    pipeline = object.__new__(JobComparisonPipeline)
    pipeline._comparison_repository = object()
    pipeline._report_repository = object()
    pipeline._db_session = object()
    comparisons = []
    reports = []

    class _UnitOfWork:
        entered = 0
        exited = 0

        def __init__(self, session):
            self.session = session

        def __enter__(self):
            type(self).entered += 1
            return self

        def __exit__(self, exc_type, exc, traceback):
            type(self).exited += 1
            return False

    def _comparison(**kwargs):
        comparisons.append(kwargs)
        return len(comparisons)

    def _report(**kwargs):
        reports.append(kwargs)
        return 100 + len(reports)

    monkeypatch.setattr("app.pipeline.PostgresUnitOfWork", _UnitOfWork)
    monkeypatch.setattr(pipeline, "_persist_postgres_comparison", _comparison)
    monkeypatch.setattr(pipeline, "_persist_postgres_report", _report)

    comparison_id, report_id = pipeline._persist_group_reports(
        group=group,
        selected_job_id="job-B",
        job_code="CLERK-01",
    )

    assert _UnitOfWork.entered == _UnitOfWork.exited == 1
    assert len(comparisons) == len(reports) == 3
    assert all(len(call["reference_jobs"]) == 2 for call in comparisons)
    assert comparison_id == 2
    assert report_id == 102


def test_compare_api_contract_stays_compatible_with_desktop_client(
    monkeypatch,
):
    import main

    seen = {}

    class _Pipeline:
        def compare_selected_job(self, job_code, company_code):
            seen.update(job_code=job_code, company_code=company_code)
            return SimpleNamespace(
                zip_bytes=b"PK-compatible-group-zip",
                zip_filename="CLERK-01_all_jobs_comparison.zip",
            )

    monkeypatch.setattr(main, "create_pipeline", lambda db: _Pipeline())
    main.app.dependency_overrides[main.get_db] = lambda: object()
    main.app.dependency_overrides[main.get_current_user] = lambda: object()

    try:
        response = TestClient(main.app).post(
            "/api/jobs/compare",
            json={"job_code": "clerk-01", "company_code": "a"},
        )
    finally:
        main.app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert "CLERK-01_all_jobs_comparison.zip" in response.headers[
        "content-disposition"
    ]
    assert seen == {"job_code": "CLERK-01", "company_code": "A"}


def test_single_job_group_returns_clean_no_benchmark_report():
    group = GroupComparisonEngine().run([_three_jobs()[0]])
    report = group.reports[0]
    assert report.reference_jobs_count == 0
    assert report.findings == []
    assert report.confidence == 0.0
